# -*- coding: utf-8 -*-
"""Persistent, sequential transcription queue. No GUI-blocking inference happens here."""
from __future__ import annotations

import json
import logging
import os
import shutil
import time
import math
import struct
import subprocess
import tempfile
import threading
import hashlib
import uuid
from datetime import datetime
from pathlib import Path

from PySide6.QtCore import Qt, QTimer, QUrl, Signal, QThread, QSize
from PySide6.QtGui import QDesktopServices, QPainter, QPen, QColor, QBrush, QPainterPath
from PySide6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QPushButton, QFrame, QSplitter, QComboBox,
    QTableWidget, QTableWidgetItem, QTextEdit, QFileDialog, QMessageBox,
    QHeaderView, QProgressBar, QAbstractItemView, QSizePolicy, QSlider,
)

from app_config import atomic_write_json
from subtitle_export import render_subtitles as _render_subtitles

SUPPORTED_EXTENSIONS = {
    ".wav", ".mp3", ".m4a", ".mp4", ".avi", ".mov", ".mkv", ".flac",
    ".aac", ".wma", ".wmv", ".alac", ".ogg", ".aiff", ".webm", ".flv",
    ".mpeg", ".mpg", ".ts",
}
STATUS_LABELS = {
    "queued": "В очереди",
    "processing": "Обрабатывается",
    "completed": "Готово",
    "failed": "Ошибка",
    "cancelled": "Отменено",
    "paused": "Приостановлено",
}

def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")

def _format_duration(seconds: float) -> str:
    seconds = max(0, int(seconds))
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}" if hours else f"{minutes:02d}:{secs:02d}"



class MediaQueue:
    """Small data helper for the persistent queue; source files are never modified or deleted."""
    @staticmethod
    def path_key(raw_path: str) -> str:
        return os.path.normcase(os.path.realpath(os.path.expanduser(raw_path)))

    @staticmethod
    def file_facts(path: Path) -> dict:
        try:
            stat = path.stat()
            size = int(stat.st_size)
        except OSError:
            size = 0
        return {
            "file_name": path.name,
            "format": path.suffix.upper().lstrip(".") or "—",
            "file_size": size,
        }

    @staticmethod
    def create_task(path: Path) -> dict:
        return {
            "id": uuid.uuid4().hex,
            "path": str(path),
            **MediaQueue.file_facts(path),
            "duration_ms": 0,
            "audio_stream_available": None,
            "waveform_status": "queued",
            "waveform_cache_key": "",
            "waveform_error": "",
            "status": "queued",
            "added_at": _now(),
            "started_at": "",
            "started_epoch": 0,
            "finished_at": "",
            "elapsed": 0,
            "progress": 0,
            "partial_text": "",
            "resume_seconds": 0.0,
            "result_path": "",
            "error": "",
        }


def _format_size(size_bytes: int) -> str:
    amount = max(0, int(size_bytes or 0))
    if amount < 1024:
        return f"{amount} Б"
    if amount < 1024 * 1024:
        return f"{amount / 1024:.0f} КБ"
    if amount < 1024 ** 3:
        return f"{amount / (1024 ** 2):.1f} МБ"
    return f"{amount / (1024 ** 3):.2f} ГБ"


class AudioEngine:
    """Small adapter around Qt Multimedia; all playback state remains owned by QMediaPlayer."""
    def __init__(self, player, audio_output):
        self.player = player
        self.audio_output = audio_output

    def load(self, path: str | Path):
        self.player.stop()
        self.player.setSource(QUrl.fromLocalFile(str(Path(path).resolve())) if path else QUrl())

    def play_pause(self):
        state = self.player.playbackState()
        if state == QMediaPlayer.PlaybackState.PlayingState:
            self.player.pause()
        else:
            duration = max(0, int(self.player.duration()))
            if duration and int(self.player.position()) >= duration:
                self.player.setPosition(0)
            self.player.play()

    def stop(self):
        self.player.stop()
        if int(self.player.duration()) > 0:
            try:
                self.player.setPosition(0)
            except RuntimeError:
                pass

    def seek(self, position_ms: int) -> bool:
        if not self.player.isSeekable():
            return False
        duration = max(0, int(self.player.duration()))
        self.player.setPosition(max(0, min(int(position_ms), duration)) if duration else max(0, int(position_ms)))
        return True

    def step(self, delta_ms: int) -> bool:
        if not self.player.isSeekable():
            return False
        return self.seek(self.player.position() + int(delta_ms))

    def set_volume(self, value: int):
        self.audio_output.setVolume(max(0, min(100, int(value))) / 100.0)

    def set_rate(self, value: float) -> bool:
        try:
            self.player.setPlaybackRate(float(value))
            return True
        except (AttributeError, RuntimeError, TypeError):
            return False


class WaveformWidget(QWidget):
    """Real amplitude data with click/drag seeking and a playback-position cursor."""
    seekRequested = Signal(int)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.samples: list[float] = []
        self.position_ms = 0
        self.duration_ms = 0
        self.loading = False
        self._dragging = False
        self.setMinimumHeight(88)
        self.setMaximumHeight(140)
        self.setMouseTracking(True)
        self.setObjectName("audioWaveform")
        self.setToolTip("Нажмите на волну или перетащите курсор, чтобы перейти к нужному времени.")

    def sizeHint(self):
        return QSize(560, 104)

    def set_samples(self, samples: list[float], duration_ms: int | None = None):
        self.samples = list(samples or [])
        if duration_ms is not None:
            self.duration_ms = max(0, int(duration_ms))
        self.loading = False
        self.update()

    def set_loading(self, loading: bool):
        self.loading = bool(loading)
        self.update()

    def set_playhead(self, position_ms: int, duration_ms: int | None = None):
        self.position_ms = max(0, int(position_ms))
        if duration_ms is not None and duration_ms > 0:
            self.duration_ms = int(duration_ms)
        self.update()

    def clear(self):
        self.samples = []
        self.position_ms = 0
        self.duration_ms = 0
        self.loading = False
        self.update()

    def _position_from_x(self, x: float) -> int | None:
        if self.duration_ms <= 0 or self.width() <= 8:
            return None
        usable = max(1.0, float(self.width() - 16))
        fraction = max(0.0, min(1.0, (float(x) - 8.0) / usable))
        return int(fraction * self.duration_ms)

    def mousePressEvent(self, event):
        if event.button() == Qt.MouseButton.LeftButton and self.duration_ms > 0:
            self._dragging = True
            position = self._position_from_x(event.position().x())
            if position is not None:
                self.position_ms = position
                self.update()
                self.seekRequested.emit(position)
            event.accept()
            return
        super().mousePressEvent(event)

    def mouseMoveEvent(self, event):
        if self._dragging:
            position = self._position_from_x(event.position().x())
            if position is not None:
                self.position_ms = position
                self.update()
                self.seekRequested.emit(position)
            event.accept()
            return
        super().mouseMoveEvent(event)

    def mouseReleaseEvent(self, event):
        if self._dragging and event.button() == Qt.MouseButton.LeftButton:
            self._dragging = False
            position = self._position_from_x(event.position().x())
            if position is not None:
                self.position_ms = position
                self.update()
                self.seekRequested.emit(position)
            event.accept()
            return
        super().mouseReleaseEvent(event)

    def paintEvent(self, event):
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing, True)
        rect = self.rect().adjusted(3, 4, -3, -4)
        path = QPainterPath()
        path.addRoundedRect(rect, 10, 10)
        painter.fillPath(path, QBrush(QColor(5, 21, 56, 130)))
        painter.setPen(QPen(QColor(89, 145, 204, 80), 1))
        painter.drawPath(path)
        inner = rect.adjusted(5, 6, -5, -6)
        mid = inner.center().y()
        painter.setPen(QPen(QColor(115, 159, 205, 85), 1))
        painter.drawLine(int(inner.left()), int(mid), int(inner.right()), int(mid))

        if not self.samples:
            painter.setPen(QColor(155, 183, 218, 200))
            message = "Подготавливаем аудиодорожку…" if self.loading else "Волновая форма ещё не рассчитана"
            painter.drawText(inner, Qt.AlignmentFlag.AlignCenter, message)
        else:
            count = len(self.samples)
            max_amp = max(self.samples) or 1.0
            threshold = max(0.006, max_amp * 0.06)
            usable_width = max(1.0, float(inner.width()))
            played_x = inner.left() + usable_width * (
                min(1.0, self.position_ms / self.duration_ms) if self.duration_ms else 0.0
            )
            painter.save()
            painter.setClipRect(inner)
            bar_width = usable_width / max(1, count)
            for index, value in enumerate(self.samples):
                normalized = min(1.0, max(0.0, value / max_amp))
                amplitude = max(1.5, normalized * inner.height() * 0.46)
                x = inner.left() + (index + 0.5) * bar_width
                is_played = x <= played_x
                if value < threshold:
                    color = QColor(105, 134, 171, 120 if not is_played else 165)
                elif is_played:
                    color = QColor(144, 225, 255, 245)
                else:
                    color = QColor(35, 151, 232, 220)
                painter.setPen(QPen(color, max(1.0, min(4.0, bar_width * 0.68)),
                                    Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap))
                painter.drawLine(int(x), int(mid - amplitude), int(x), int(mid + amplitude))
            painter.restore()

        if self.duration_ms > 0:
            usable_width = max(1.0, float(inner.width()))
            x = inner.left() + usable_width * min(1.0, self.position_ms / self.duration_ms)
            painter.setPen(QPen(QColor(242, 250, 255, 240), 1.5))
            painter.drawLine(int(x), int(inner.top()), int(x), int(inner.bottom()))
        if self.loading and self.samples:
            painter.setPen(QColor(220, 241, 255, 200))
            painter.drawText(rect.adjusted(8, 4, -8, -4), Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignTop,
                             "Обновление…")
        painter.end()


class WaveformExtractor(QThread):
    """Streaming FFmpeg PCM extractor with file-fingerprint cache; never buffers the whole audio."""
    audio_stream_detected = Signal(str)
    completed = Signal(str, object, int, str)  # path, amplitudes, duration_ms, cache_key
    failed = Signal(str, str)

    SAMPLE_RATE = 8000
    BUCKET_MS = 100
    MAX_BARS = 1800

    def __init__(self, source_path: str, cache_dir: Path, parent=None, force: bool = False):
        super().__init__(parent)
        self.source_path = str(Path(source_path).resolve())
        self.cache_dir = Path(cache_dir)
        self.force = bool(force)
        self._cancel_event = threading.Event()
        self._process = None
        self._process_guard = threading.RLock()

    def cancel(self):
        self._cancel_event.set()
        with self._process_guard:
            process = self._process
        if process is not None and process.poll() is None:
            try:
                process.terminate()
            except OSError:
                pass

    def _fingerprint(self) -> tuple[str, Path, int]:
        source = Path(self.source_path)
        stat = source.stat()
        identity = {
            "path": str(source.resolve()),
            "size": int(stat.st_size),
            "mtime_ns": int(stat.st_mtime_ns),
            "sample_rate": self.SAMPLE_RATE,
            "bucket_ms": self.BUCKET_MS,
            "version": 1,
        }
        digest = hashlib.sha256(json.dumps(identity, sort_keys=True).encode("utf-8")).hexdigest()
        return digest, self.cache_dir / f"{digest}.json", int(stat.st_size)

    @staticmethod
    def _rms(data: bytes) -> float:
        if len(data) < 2:
            return 0.0
        usable = len(data) - (len(data) % 2)
        count = usable // 2
        samples = struct.unpack("<" + "h" * count, data[:usable])
        return math.sqrt(sum(value * value for value in samples) / count) / 32768.0

    def _emit_cached(self, cache_path: Path, cache_key: str) -> bool:
        try:
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
            if payload.get("cache_key") != cache_key or payload.get("version") != 1:
                return False
            bars = payload.get("samples")
            duration_ms = int(payload.get("duration_ms", 0))
            if not isinstance(bars, list) or not bars or duration_ms <= 0:
                return False
            clean = [max(0.0, min(1.0, float(value))) for value in bars]
            self.audio_stream_detected.emit(self.source_path)
            self.completed.emit(self.source_path, clean, duration_ms, cache_key)
            return True
        except (OSError, ValueError, TypeError, json.JSONDecodeError):
            return False

    def run(self):
        process = None
        try:
            source = Path(self.source_path)
            if not source.is_file():
                raise FileNotFoundError(f"Файл не найден: {source}")
            if source.stat().st_size <= 0:
                raise ValueError("Файл пустой.")
            cache_key, cache_path, size = self._fingerprint()
            if not self.force and self._emit_cached(cache_path, cache_key):
                return

            import imageio_ffmpeg
            ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
            cmd = [
                ffmpeg, "-hide_banner", "-loglevel", "error", "-nostdin",
                "-i", self.source_path, "-map", "0:a:0", "-vn", "-ac", "1",
                "-ar", str(self.SAMPLE_RATE), "-f", "s16le", "pipe:1",
            ]
            create_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0) if os.name == "nt" else 0
            bars: list[float] = []
            sample_count = 0
            stream_announced = False
            sample_bytes = self.SAMPLE_RATE * 2 * self.BUCKET_MS // 1000
            pending = bytearray()
            with tempfile.TemporaryFile(mode="w+b") as err_stream:
                process = subprocess.Popen(
                    cmd, stdout=subprocess.PIPE, stderr=err_stream,
                    stdin=subprocess.DEVNULL, creationflags=create_flags,
                )
                with self._process_guard:
                    self._process = process
                assert process.stdout is not None
                while True:
                    if self._cancel_event.is_set():
                        process.terminate()
                        try:
                            process.wait(timeout=2)
                        except subprocess.TimeoutExpired:
                            process.kill()
                            process.wait()
                        raise InterruptedError("Анализ аудиодорожки отменён.")
                    # read1() returns whatever FFmpeg has produced in one underlying read.
                    # Buffered read(64 KiB) could delay stream detection by several seconds.
                    data = process.stdout.read1(8192) if hasattr(process.stdout, "read1") else process.stdout.read(8192)
                    if not data:
                        break
                    if not stream_announced:
                        self.audio_stream_detected.emit(self.source_path)
                        stream_announced = True
                    pending.extend(data)
                    # Fixed 100 ms slices give a stable timeline independent of pipe read sizes.
                    while len(pending) >= sample_bytes:
                        block = bytes(pending[:sample_bytes])
                        del pending[:sample_bytes]
                        bars.append(self._rms(block))
                        sample_count += sample_bytes // 2
                if pending:
                    usable = len(pending) - (len(pending) % 2)
                    if usable:
                        bars.append(self._rms(bytes(pending[:usable])))
                        sample_count += usable // 2
                code = process.wait()
                err_stream.seek(0)
                stderr_data = err_stream.read()
            with self._process_guard:
                self._process = None
            if self._cancel_event.is_set():
                return
            if code != 0:
                detail = stderr_data.decode("utf-8", errors="replace").strip()
                if "matches no streams" in detail.lower() or "does not contain any stream" in detail.lower():
                    raise RuntimeError("В видео не найдена аудиодорожка.")
                raise RuntimeError(detail[-1500:] or f"FFmpeg завершился с кодом {code}.")
            if not bars or sample_count <= 0:
                raise RuntimeError("Не удалось извлечь аудиодорожку. Файл может не содержать звук.")
            duration_ms = int(sample_count * 1000 / self.SAMPLE_RATE)
            if len(bars) > self.MAX_BARS:
                source_bars = bars
                step = len(source_bars) / self.MAX_BARS
                bars = [
                    max(source_bars[int(i * step):max(int(i * step) + 1, int((i + 1) * step))], default=0.0)
                    for i in range(self.MAX_BARS)
                ]
            self.cache_dir.mkdir(parents=True, exist_ok=True)
            payload = {"version": 1, "cache_key": cache_key, "duration_ms": duration_ms, "samples": bars}
            tmp = cache_path.with_suffix(".tmp")
            tmp.write_text(json.dumps(payload, separators=(",", ":")), encoding="utf-8")
            os.replace(tmp, cache_path)
            logging.getLogger("processing.audio").info(
                "Waveform extracted: file=%s bytes=%d duration_ms=%d samples=%d cache=%s",
                source.name, size, duration_ms, len(bars), cache_path.name
            )
            self.completed.emit(self.source_path, bars, duration_ms, cache_key)
        except InterruptedError:
            return
        except Exception as exc:
            if process is not None and process.poll() is None:
                try:
                    process.terminate()
                    process.wait(timeout=2)
                except Exception:
                    try:
                        process.kill()
                        process.wait()
                    except Exception:
                        pass
            with self._process_guard:
                self._process = None
            logging.getLogger("processing.audio").warning("Waveform extraction failed for %s: %s",
                                                          self.source_path, exc)
            if not self._cancel_event.is_set():
                self.failed.emit(self.source_path, str(exc))


# Backwards compatible name used by the first player implementation.
WaveformAnalysisWorker = WaveformExtractor


class PlayerControls(QWidget):
    """Reusable transport, timeline, volume and playback-rate controls."""
    playPauseClicked = Signal()
    stopClicked = Signal()
    back10Clicked = Signal()
    forward10Clicked = Signal()
    seekRequested = Signal(int)
    volumeChanged = Signal(int)
    rateChanged = Signal(float)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setObjectName("playerControls")
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 4, 0, 0)
        root.setSpacing(7)

        row = QHBoxLayout()
        row.setSpacing(7)
        self.back10_button = QPushButton("−10 с")
        self.back10_button.setProperty("role", "tertiary")
        self.back10_button.setToolTip("Перемотать на 10 секунд назад")
        self.play_button = QPushButton("▶  Воспроизвести")
        self.play_button.setProperty("role", "primary")
        self.stop_button = QPushButton("■  Стоп")
        self.stop_button.setProperty("role", "tertiary")
        self.forward10_button = QPushButton("+10 с")
        self.forward10_button.setProperty("role", "tertiary")
        self.jump_combo = QComboBox()
        self.jump_combo.setObjectName("playerJumpCombo")
        for seconds, label in ((5, "5 с"), (10, "10 с"), (15, "15 с"), (30, "30 с")):
            self.jump_combo.addItem(label, seconds)
        self.jump_combo.setCurrentIndex(1)
        self.jump_seconds = 10
        row.addWidget(self.back10_button)
        row.addWidget(self.play_button, 1)
        row.addWidget(self.stop_button)
        row.addWidget(self.forward10_button)
        row.addWidget(self.jump_combo)

        self.time_label = QLabel("00:00 / 00:00")
        self.time_label.setObjectName("playerTime")
        self.time_label.setMinimumWidth(104)
        self.time_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        row.addWidget(self.time_label)

        root.addLayout(row)

        timeline_row = QHBoxLayout()
        timeline_row.setSpacing(8)
        self.seek_slider = QSlider(Qt.Orientation.Horizontal)
        self.seek_slider.setObjectName("playerSeekSlider")
        self.seek_slider.setRange(0, 0)
        timeline_row.addWidget(self.seek_slider, 1)

        volume_label = QLabel("Громкость")
        volume_label.setObjectName("mutedLabel")
        timeline_row.addWidget(volume_label)
        self.volume_slider = QSlider(Qt.Orientation.Horizontal)
        self.volume_slider.setObjectName("playerVolumeSlider")
        self.volume_slider.setRange(0, 100)
        self.volume_slider.setValue(78)
        self.volume_slider.setMaximumWidth(100)
        self.volume_slider.setToolTip("Громкость воспроизведения")
        timeline_row.addWidget(self.volume_slider)

        speed_label = QLabel("Скорость")
        speed_label.setObjectName("mutedLabel")
        timeline_row.addWidget(speed_label)
        self.speed_combo = QComboBox()
        self.speed_combo.setObjectName("playerSpeedCombo")
        for rate, label in ((0.5, "0,5×"), (0.75, "0,75×"), (1.0, "1×"), (1.25, "1,25×"),
                            (1.5, "1,5×"), (1.75, "1,75×"), (2.0, "2×")):
            self.speed_combo.addItem(label, rate)
        self.speed_combo.setCurrentIndex(2)
        self.speed_combo.setMaximumWidth(90)
        timeline_row.addWidget(self.speed_combo)
        root.addLayout(timeline_row)

        self.back10_button.clicked.connect(self.back10Clicked)
        self.play_button.clicked.connect(self.playPauseClicked)
        self.stop_button.clicked.connect(self.stopClicked)
        self.forward10_button.clicked.connect(self.forward10Clicked)
        self.jump_combo.currentIndexChanged.connect(self._jump_changed)
        self.seek_slider.sliderMoved.connect(self.seekRequested)
        self.volume_slider.valueChanged.connect(self.volumeChanged)
        self.speed_combo.currentIndexChanged.connect(lambda _: self.rateChanged.emit(float(self.speed_combo.currentData())))

    def _jump_changed(self, _index):
        self.jump_seconds = int(self.jump_combo.currentData() or 10)
        self.back10_button.setText(f"−{self.jump_seconds} с")
        self.forward10_button.setText(f"+{self.jump_seconds} с")

    def set_enabled(self, enabled: bool):
        for widget in (
            self.back10_button, self.play_button, self.stop_button, self.forward10_button,
            self.seek_slider, self.volume_slider, self.speed_combo, self.jump_combo,
        ):
            widget.setEnabled(enabled)

    def set_playing(self, playing: bool):
        self.play_button.setText("Ⅱ  Пауза" if playing else "▶  Воспроизвести")


class QueuePage(QWidget):
    """Persistent work queue with one active model inference at a time."""
    ai_requested = Signal(str, str)

    def __init__(self, project_root: Path, worker_factory, on_back, parent=None):
        super().__init__(parent)
        # Load Qt Multimedia only when the user opens the file queue; keep the home screen launch light.
        global QMediaPlayer, QAudioOutput
        from PySide6.QtMultimedia import QMediaPlayer, QAudioOutput

        self.project_root = Path(project_root)
        self.worker_factory = worker_factory
        self.on_back = on_back
        self.data_dir = self.project_root / "data"
        self.queue_file = self.data_dir / "queue.json"
        self.results_dir = self.project_root / "results"
        self.results_dir.mkdir(parents=True, exist_ok=True)
        self.data_dir.mkdir(parents=True, exist_ok=True)
        self.tasks: list[dict] = []
        self.active_worker = None
        self.active_task_id = None
        self.queue_running = False
        self.cancel_requested = False
        self.stop_requested = False
        self.discard_requested = False
        self.shutdown_requested = False
        self._displayed_task_id = None
        self._loading_editor = False
        self.busy_changed = None
        self.external_busy_check = None
        self.setObjectName("queuePage")
        self.setAcceptDrops(True)

        root = QVBoxLayout(self)
        root.setContentsMargins(28, 22, 28, 24)
        root.setSpacing(14)

        header = QHBoxLayout()
        self.back_button = QPushButton("←  Назад")
        self.back_button.setObjectName("pushButton_back_to_main")
        self.back_button.setProperty("role", "back")
        self.back_button.setMinimumHeight(40)
        self.back_button.clicked.connect(self.on_back)
        header.addWidget(self.back_button, 0, Qt.AlignmentFlag.AlignTop)
        heading = QVBoxLayout()
        heading.setSpacing(3)
        title = QLabel("Очередь обработки")
        title.setObjectName("dialogTitle")
        subtitle = QLabel("Добавляйте несколько файлов. СПИКЕР обрабатывает их последовательно, а очередь сохраняется между запусками.")
        subtitle.setObjectName("mutedLabel")
        subtitle.setWordWrap(True)
        heading.addWidget(title)
        heading.addWidget(subtitle)
        header.addLayout(heading, 1)
        root.addLayout(header)

        self.media_splitter = QSplitter(Qt.Orientation.Horizontal)
        self.media_splitter.setObjectName("mediaWorkspaceSplitter")
        self.media_splitter.setChildrenCollapsible(False)

        # Left: persistent media/transcription queue.
        self.queue_panel = QFrame()
        self.queue_panel.setObjectName("panel")
        left_layout = QVBoxLayout(self.queue_panel)
        left_layout.setContentsMargins(12, 12, 12, 12)
        left_layout.setSpacing(9)

        queue_heading = QLabel("ФАЙЛЫ")
        queue_heading.setObjectName("sectionLabel")
        left_layout.addWidget(queue_heading)

        controls_grid = QGridLayout()
        controls_grid.setHorizontalSpacing(6)
        controls_grid.setVerticalSpacing(6)
        self.add_button = QPushButton("＋ Добавить")
        self.add_button.setProperty("role", "primary")
        self.add_button.clicked.connect(self.add_files_dialog)
        controls_grid.addWidget(self.add_button, 0, 0)
        self.start_button = QPushButton("▶ Очередь")
        self.start_button.setProperty("role", "secondary")
        self.start_button.setToolTip("Запустить новые задачи и продолжить приостановленные с сохранённого сегмента.")
        self.start_button.clicked.connect(self.start_queue)
        controls_grid.addWidget(self.start_button, 0, 1)
        self.pause_button = QPushButton("Пауза")
        self.pause_button.setProperty("role", "tertiary")
        self.pause_button.setToolTip("Приостановить текущую задачу на ближайшей границе сегмента, не сбрасывая прогресс.")
        self.pause_button.clicked.connect(self.toggle_pause)
        self.pause_button.setEnabled(False)
        controls_grid.addWidget(self.pause_button, 1, 0)
        self.cancel_button = QPushButton("Остановить и сохранить")
        self.cancel_button.setProperty("role", "tertiary")
        self.cancel_button.setToolTip("Остановить текущую задачу, сохранив текст и позицию для продолжения.")
        self.cancel_button.clicked.connect(self.stop_current_task)
        self.cancel_button.setEnabled(False)
        controls_grid.addWidget(self.cancel_button, 1, 1)
        self.cancel_task_button = QPushButton("Отменить задачу")
        self.cancel_task_button.setProperty("role", "tertiary")
        self.cancel_task_button.setToolTip("Отменить выбранную задачу и сбросить её контрольную точку. Исходный файл не удаляется.")
        self.cancel_task_button.clicked.connect(self.cancel_selected_task)
        controls_grid.addWidget(self.cancel_task_button, 2, 0)
        self.retry_button = QPushButton("Продолжить / повторить")
        self.retry_button.setProperty("role", "tertiary")
        self.retry_button.clicked.connect(self.retry_selected)
        controls_grid.addWidget(self.retry_button, 2, 1)
        self.remove_button = QPushButton("Удалить из списка")
        self.remove_button.setProperty("role", "tertiary")
        self.remove_button.setToolTip("Удалить только запись из очереди; исходный аудио-/видеофайл останется на диске.")
        self.remove_button.clicked.connect(self.remove_selected)
        controls_grid.addWidget(self.remove_button, 3, 0, 1, 2)
        left_layout.addLayout(controls_grid)

        self.table = QTableWidget(0, 5)
        self.table.setObjectName("queueTable")
        self.table.setHorizontalHeaderLabels(["Файл", "Длит.", "Размер", "Тип", "Статус"])
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.setAlternatingRowColors(True)
        self.table.verticalHeader().setVisible(False)
        header_view = self.table.horizontalHeader()
        header_view.setSectionResizeMode(0, QHeaderView.ResizeMode.Stretch)
        for column in range(1, 5):
            header_view.setSectionResizeMode(column, QHeaderView.ResizeMode.ResizeToContents)
        self.table.setMinimumHeight(240)
        self.table.itemSelectionChanged.connect(self._selection_changed)
        self.table.cellDoubleClicked.connect(self.open_selected_source)
        left_layout.addWidget(self.table, 1)
        self.media_splitter.addWidget(self.queue_panel)

        # Right: selected media, actual waveform, transport controls, processing state and text.
        self.player_panel = QFrame()
        self.player_panel.setObjectName("panel")
        player_layout = QVBoxLayout(self.player_panel)
        player_layout.setContentsMargins(16, 14, 16, 14)
        player_layout.setSpacing(10)

        title_row = QHBoxLayout()
        title_stack = QVBoxLayout()
        title_stack.setSpacing(3)
        self.player_title = QLabel("Выберите файл")
        self.player_title.setObjectName("dialogTitle")
        self.player_title.setWordWrap(True)
        self.player_metadata = QLabel("Длительность · Размер · Формат")
        self.player_metadata.setObjectName("mutedLabel")
        self.player_metadata.setWordWrap(True)
        title_stack.addWidget(self.player_title)
        title_stack.addWidget(self.player_metadata)
        title_row.addLayout(title_stack, 1)
        self.waveform_button = QPushButton("Обновить волну")
        self.waveform_button.setProperty("role", "tertiary")
        self.waveform_button.setEnabled(False)
        self.waveform_button.setToolTip("Повторно рассчитать амплитудную форму сигнала")
        self.waveform_button.clicked.connect(lambda: self.build_waveform(force=True))
        title_row.addWidget(self.waveform_button, 0, Qt.AlignmentFlag.AlignTop)
        player_layout.addLayout(title_row)

        self.waveform = WaveformWidget()
        self.waveform.seekRequested.connect(self._seek_preview)
        player_layout.addWidget(self.waveform, 1)

        self.player_controls = PlayerControls()
        self.player_controls.set_enabled(False)
        self.player_controls.playPauseClicked.connect(self.toggle_preview)
        self.player_controls.stopClicked.connect(self.stop_preview)
        self.player_controls.back10Clicked.connect(
            lambda: self._step_preview(-self.player_controls.jump_seconds * 1000)
        )
        self.player_controls.forward10Clicked.connect(
            lambda: self._step_preview(self.player_controls.jump_seconds * 1000)
        )
        self.player_controls.seekRequested.connect(self._seek_preview)
        self.player_controls.volumeChanged.connect(self._set_volume)
        self.player_controls.rateChanged.connect(self._set_playback_rate)
        player_layout.addWidget(self.player_controls)

        # Backwards-compatible aliases used by current slots and status handlers.
        self.play_button = self.player_controls.play_button
        self.stop_preview_button = self.player_controls.stop_button
        self.seek_slider = self.player_controls.seek_slider
        self.audio_time_label = self.player_controls.time_label
        self.volume_slider = self.player_controls.volume_slider
        self.speed_combo = self.player_controls.speed_combo
        self.audio_output = QAudioOutput(self)
        self.audio_output.setVolume(0.78)
        self.media_player = QMediaPlayer(self)
        self.media_player.setAudioOutput(self.audio_output)
        self.audio_engine = AudioEngine(self.media_player, self.audio_output)
        self.media_player.positionChanged.connect(self._preview_position_changed)
        self.media_player.durationChanged.connect(self._preview_duration_changed)
        self.media_player.playbackStateChanged.connect(self._preview_state_changed)
        self.media_player.mediaStatusChanged.connect(self._preview_media_status_changed)
        self.media_player.seekableChanged.connect(self._preview_seekable_changed)
        self.media_player.errorOccurred.connect(self._preview_error)
        self._preview_source = ""
        self._player_error_sources: set[str] = set()
        self._waveform_worker = None
        self._waveform_task_id = None
        self._waveform_pending_ids: list[str] = []
        self._waveform_samples_by_path: dict[str, list[float]] = {}
        self._waveform_force_by_id: set[str] = set()
        self._waveform_preempted_ids: set[str] = set()
        self._waveform_cache_dir = self.data_dir / "waveform_cache"
        self._waveform_cache_dir.mkdir(parents=True, exist_ok=True)

        self.audio_status = QLabel("Добавьте файл. После проверки аудиодорожки волна появится здесь.")
        self.audio_status.setObjectName("mutedLabel")
        self.audio_status.setWordWrap(True)
        player_layout.addWidget(self.audio_status)

        activity = QFrame()
        activity.setObjectName("panel")
        activity_layout = QVBoxLayout(activity)
        activity_layout.setContentsMargins(12, 8, 12, 8)
        activity_layout.setSpacing(6)
        status_row = QHBoxLayout()
        self.current_label = QLabel("Распознавание не запущено")
        self.current_label.setObjectName("cardTitle")
        status_row.addWidget(self.current_label, 1)
        self.elapsed_label = QLabel("")
        self.elapsed_label.setObjectName("mutedLabel")
        status_row.addWidget(self.elapsed_label)
        activity_layout.addLayout(status_row)
        self.progress_bar = QProgressBar()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(0)
        self.progress_bar.setTextVisible(False)
        activity_layout.addWidget(self.progress_bar)
        self.progress_message = QLabel("Выберите файл в очереди или добавьте новые.")
        self.progress_message.setObjectName("mutedLabel")
        self.progress_message.setWordWrap(True)
        activity_layout.addWidget(self.progress_message)
        player_layout.addWidget(activity, 0)

        result_panel = QFrame()
        result_panel.setObjectName("panel")
        result_layout = QVBoxLayout(result_panel)
        result_layout.setContentsMargins(12, 8, 12, 8)
        result_layout.setSpacing(6)
        result_header = QHBoxLayout()
        result_title = QLabel("Результат транскрибации")
        result_title.setObjectName("sectionLabel")
        result_header.addWidget(result_title)
        result_header.addStretch(1)
        self.ai_result_button = QPushButton("Обработать с ИИ")
        self.ai_result_button.setProperty("role", "secondary")
        self.ai_result_button.setToolTip("Передать готовую расшифровку в отдельную панель локального ИИ или API")
        self.ai_result_button.clicked.connect(self.send_result_to_ai)
        result_header.addWidget(self.ai_result_button)
        self.save_result_button = QPushButton("Экспорт…")
        self.save_result_button.setProperty("role", "secondary")
        self.save_result_button.clicked.connect(self.save_current_result_as)
        result_header.addWidget(self.save_result_button)
        result_layout.addLayout(result_header)
        self.result_editor = QTextEdit()
        self.result_editor.setObjectName("queueResultEditor")
        self.result_editor.setPlaceholderText("Здесь появится результат выбранного файла. Во время распознавания показываются уже готовые фрагменты.")
        self.result_editor.setMinimumHeight(92)
        self.result_editor.setReadOnly(True)
        result_layout.addWidget(self.result_editor)
        self._editor_save_timer = QTimer(self)
        self._editor_save_timer.setSingleShot(True)
        self._editor_save_timer.setInterval(750)
        self._editor_save_timer.timeout.connect(self._persist_result_editor)
        self.result_editor.textChanged.connect(self._result_editor_text_changed)
        player_layout.addWidget(result_panel, 2)

        self.media_splitter.addWidget(self.player_panel)
        self.media_splitter.setStretchFactor(0, 4)
        self.media_splitter.setStretchFactor(1, 6)
        self.media_splitter.setSizes([390, 610])
        root.addWidget(self.media_splitter, 1)

        # Compatibility aliases from previous audio preview layout.
        self.preview_panel = self.player_panel
        

        self._clock = QTimer(self)
        self._clock.setInterval(1000)
        self._clock.timeout.connect(self._update_clock)
        self._clock.start()
        self._load_queue()

    def _load_queue(self):
        changed = False
        try:
            payload = json.loads(self.queue_file.read_text(encoding="utf-8"))
            if isinstance(payload, dict) and isinstance(payload.get("tasks"), list):
                self.tasks = [task for task in payload["tasks"] if isinstance(task, dict) and task.get("path")]
        except json.JSONDecodeError:
            self.tasks = []
            changed = True
            # Preserve a broken queue snapshot for troubleshooting before starting a new queue.
            try:
                backup = self.queue_file.with_name(f"queue.corrupt.{int(time.time())}.json")
                shutil.copy2(self.queue_file, backup)
                logging.getLogger("app").exception("Queue JSON was invalid; copied to %s", backup)
            except OSError:
                logging.getLogger("app").exception("Queue JSON was invalid and couldn't be backed up.")
        except OSError:
            self.tasks = []

        for task in self.tasks:
            task.setdefault("id", uuid.uuid4().hex)
            task.setdefault("added_at", _now())
            task.setdefault("progress", 0)
            task.setdefault("elapsed", 0)
            task.setdefault("partial_text", "")
            if not isinstance(task.get("segments"), list):
                task["segments"] = []
            task.setdefault("resume_seconds", 0.0)
            try:
                task["resume_seconds"] = max(0.0, float(task.get("resume_seconds", 0) or 0))
            except (TypeError, ValueError):
                task["resume_seconds"] = 0.0
                changed = True
            task.setdefault("result_path", "")
            task.setdefault("error", "")
            try:
                path = Path(task.get("path", "")).expanduser().resolve()
                task.update({k: v for k, v in MediaQueue.file_facts(path).items() if not task.get(k)})
                if not path.is_file():
                    task["waveform_status"] = "failed"
                    task["audio_stream_available"] = False
                    task["waveform_error"] = "Исходный файл больше не найден."
            except (OSError, ValueError):
                pass
            task.setdefault("duration_ms", 0)
            task.setdefault("audio_stream_available", None)
            task.setdefault("waveform_status", "queued")
            task.setdefault("waveform_cache_key", "")
            task.setdefault("waveform_error", "")
            if task.get("waveform_status") in {"loading", "ready"}:
                # Reload compact waveform samples from the fingerprinted cache on demand.
                task["waveform_status"] = "queued"
                changed = True
            if task.get("waveform_status") == "queued" and task.get("id") not in self._waveform_pending_ids:
                self._waveform_pending_ids.append(task.get("id"))
            if task.get("status") == "processing":
                valid_checkpoint = bool(task.get("partial_text")) and float(task.get("resume_seconds", 0) or 0) > 0
                task["status"] = "paused" if valid_checkpoint else "queued"
                task["started_at"] = ""
                task["started_epoch"] = 0
                if not valid_checkpoint:
                    task["progress"] = 0
                    task["resume_seconds"] = 0.0
                    task["partial_text"] = ""
                task["error"] = (
                    "Предыдущий запуск прерван. Можно продолжить с последнего сохранённого сегмента."
                    if valid_checkpoint else ""
                )
                changed = True
            if task.get("status") not in STATUS_LABELS:
                task["status"] = "queued"
                changed = True
        if changed:
            self._save_queue()
        self._render_table()
        if self._waveform_pending_ids:
            QTimer.singleShot(0, self._start_next_waveform_extraction)

    def _save_queue(self):
        atomic_write_json(self.queue_file, {"version": 1, "updated_at": _now(), "tasks": self.tasks})

    def _task_by_id(self, task_id):
        return next((item for item in self.tasks if item.get("id") == task_id), None)

    def _selected_task(self):
        row = self.table.currentRow()
        if row < 0:
            return None
        item = self.table.item(row, 0)
        return self._task_by_id(item.data(Qt.ItemDataRole.UserRole)) if item else None

    def _render_table(self, select_id=None):
        if select_id is None and self.table.currentRow() >= 0:
            first = self.table.item(self.table.currentRow(), 0)
            select_id = first.data(Qt.ItemDataRole.UserRole) if first else None
        self.table.blockSignals(True)
        self.table.setRowCount(len(self.tasks))
        select_row = -1
        for row, task in enumerate(self.tasks):
            path = Path(task.get("path", ""))
            name = path.name or task.get("path", "(неизвестный файл)")
            name_item = QTableWidgetItem(name)
            name_item.setToolTip(task.get("path", ""))
            name_item.setData(Qt.ItemDataRole.UserRole, task.get("id"))
            if task.get("error") or task.get("waveform_error"):
                name_item.setToolTip(
                    f"{task.get('path', '')}\n\n{task.get('error', '')}\n{task.get('waveform_error', '')}"
                )

            duration_ms = int(task.get("duration_ms", 0) or 0)
            duration_item = QTableWidgetItem(_format_duration(duration_ms / 1000) if duration_ms else "—")
            size_item = QTableWidgetItem(_format_size(int(task.get("file_size", 0) or 0)))
            format_item = QTableWidgetItem(str(task.get("format") or path.suffix.upper().lstrip(".") or "—"))

            task_status = task.get("status", "queued")
            wave_status = task.get("waveform_status", "queued")
            if task_status == "processing":
                progress = int(task.get("progress", 0) or 0)
                status = f"Обработка {progress}%" if progress > 0 else "Обработка…"
                if self.active_worker is not None and self.active_worker.is_paused():
                    status = f"Пауза {progress}%"
            elif task_status == "completed":
                status = "Готово"
            elif task_status == "failed":
                status = "Ошибка"
            elif task_status == "cancelled":
                status = "Отменено"
            elif task_status == "paused":
                status = "Приостановлено"
            elif wave_status == "loading":
                status = "Проверка аудио…"
            elif wave_status == "queued":
                status = "Подготавливается…"
            elif wave_status == "failed":
                status = "Нет аудио / ошибка"
            else:
                status = STATUS_LABELS.get(task_status, task_status)

            status_item = QTableWidgetItem(status)
            status_item.setToolTip(task.get("waveform_error", "") or task.get("error", ""))
            status_item.setData(Qt.ItemDataRole.UserRole + 1, task.get("id"))
            for col, item in enumerate((name_item, duration_item, size_item, format_item, status_item)):
                self.table.setItem(row, col, item)
            if task.get("id") == select_id:
                select_row = row
        if select_row >= 0:
            self.table.selectRow(select_row)
        elif self.tasks and self.table.currentRow() < 0:
            self.table.selectRow(0)
        self.table.blockSignals(False)
        self._selection_changed()
        pending = sum(1 for task in self.tasks if task.get("status") == "queued")
        paused_count = sum(1 for task in self.tasks if task.get("status") == "paused")
        done = sum(1 for task in self.tasks if task.get("status") == "completed")
        active = sum(1 for task in self.tasks if task.get("status") == "processing")
        resumable_pending = sum(1 for task in self.tasks if task.get("status") in {"queued", "paused"})
        self.start_button.setEnabled(self.active_worker is None and resumable_pending > 0)
        self.cancel_button.setEnabled(self.active_worker is not None)
        self.pause_button.setEnabled(self.active_worker is not None)
        self.pause_button.setText("Продолжить" if self.active_worker is not None and self.active_worker.is_paused() else "Пауза")
        selected_task = self._selected_task()
        selected_status = (selected_task or {}).get("status")
        self.retry_button.setEnabled(bool(selected_task and selected_status in {"failed", "cancelled", "paused"}))
        self.cancel_task_button.setEnabled(bool(selected_task and selected_status in {"queued", "paused", "processing", "failed"}))
        self.remove_button.setEnabled(bool(selected_task) and selected_status != "processing")
        if self.active_worker is None and self.progress_message.text().startswith(
            ("Перетащите", "В очереди:", "Очередь завершена.")
        ):
            self.progress_message.setText(
                f"В очереди: {pending} · Приостановлено: {paused_count} · Завершено: {done} · Сейчас обрабатывается: {active}"
            )

    def add_files_dialog(self):
        files, _ = QFileDialog.getOpenFileNames(
            self, "Добавить аудио- и видеофайлы", "",
            "Медиафайлы (*.wav *.mp3 *.m4a *.mp4 *.avi *.mov *.mkv *.flac *.aac *.wma *.wmv *.ogg *.aiff *.webm *.flv *.mpeg *.mpg *.ts);;Все файлы (*.*)"
        )
        self.add_paths(files)

    def add_paths(self, paths):
        existing_paths = {MediaQueue.path_key(str(t.get("path", ""))) for t in self.tasks if t.get("path")}
        added_ids = []
        skipped = []
        duplicates = []
        for raw in paths:
            try:
                path = Path(raw).expanduser().resolve()
                if not path.is_file():
                    skipped.append(f"{path.name}: файл не найден")
                    continue
                if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    skipped.append(f"{path.name}: формат не поддерживается")
                    continue
                path_key = MediaQueue.path_key(str(path))
                if path_key in existing_paths:
                    duplicates.append(path.name)
                    continue
                task = MediaQueue.create_task(path)
                self.tasks.append(task)
                existing_paths.add(path_key)
                added_ids.append(task["id"])
                self._waveform_pending_ids.append(task["id"])
            except (OSError, ValueError) as exc:
                skipped.append(f"{Path(raw).name}: {exc}")
        self._save_queue()
        self._render_table(added_ids[0] if added_ids else None)
        if added_ids:
            QTimer.singleShot(0, self._start_next_waveform_extraction)
        if skipped:
            self.progress_message.setText("Не добавлены: " + "; ".join(skipped[:4]))
        elif duplicates and not added_ids:
            self.progress_message.setText("Эти файлы уже есть в очереди: " + ", ".join(duplicates[:5]))
        elif added_ids:
            suffix = f" Пропущено повторов: {len(duplicates)}." if duplicates else ""
            self.progress_message.setText(
                f"Добавлено файлов: {len(added_ids)}. Проверяем аудиодорожки и длительность.{suffix}"
            )
        return len(added_ids)

    def _queue_waveform_task(self, task_id, priority=True, force=False):
        task = self._task_by_id(task_id)
        if not task or not Path(task.get("path", "")).is_file():
            return
        if task.get("waveform_status") == "loading" and self._waveform_task_id == task_id:
            return
        if self._waveform_worker is not None and self._waveform_task_id == task_id and self._waveform_worker.isRunning():
            return
        if priority and self._waveform_worker is not None and self._waveform_worker.isRunning():
            running_id = self._waveform_task_id
            if running_id and running_id != task_id:
                self._waveform_preempted_ids.add(running_id)
                self._waveform_worker.cancel()
        task["waveform_status"] = "queued"
        if force:
            self._waveform_force_by_id.add(task_id)
            self._waveform_samples_by_path.pop(str(Path(task["path"]).resolve()), None)
        if task_id in self._waveform_pending_ids:
            self._waveform_pending_ids.remove(task_id)
        if priority:
            self._waveform_pending_ids.insert(0, task_id)
        else:
            self._waveform_pending_ids.append(task_id)
        self._save_queue()
        QTimer.singleShot(0, self._start_next_waveform_extraction)

    def _start_next_waveform_extraction(self):
        if self._waveform_worker is not None and self._waveform_worker.isRunning():
            return
        self._waveform_worker = None
        self._waveform_task_id = None
        while self._waveform_pending_ids:
            task_id = self._waveform_pending_ids.pop(0)
            task = self._task_by_id(task_id)
            if not task:
                continue
            source = Path(task.get("path", "")).resolve()
            if not source.is_file():
                task["waveform_status"] = "failed"
                task["audio_stream_available"] = False
                task["waveform_error"] = "Исходный файл не найден."
                self._save_queue()
                continue
            self._waveform_task_id = task_id
            task["waveform_status"] = "loading"
            task["waveform_error"] = ""
            self._save_queue()
            self._render_table()
            if str(source) == self._preview_source:
                self.waveform.set_loading(True)
                self.audio_status.setText("Подготавливаем аудиодорожку…")
            worker = WaveformExtractor(
                str(source), self._waveform_cache_dir, self,
                force=task_id in self._waveform_force_by_id,
            )
            self._waveform_force_by_id.discard(task_id)
            worker.audio_stream_detected.connect(self._audio_stream_detected)
            worker.completed.connect(self._waveform_ready)
            worker.failed.connect(self._waveform_failed)
            worker.finished.connect(self._waveform_worker_finished)
            self._waveform_worker = worker
            worker.start()
            return
        self._save_queue()
        self._render_table()


    def start_queue(self):
        if self.active_worker is not None:
            return
        if not any(task.get("status") in {"queued", "paused"} for task in self.tasks):
            self.progress_message.setText("Нет файлов, ожидающих обработки или продолжения.")
            return
        self.queue_running = True
        self.cancel_requested = False
        self.stop_requested = False
        self.discard_requested = False
        self._start_next()

    def _start_next(self):
        if not self.queue_running or self.active_worker is not None:
            return
        task = next((item for item in self.tasks if item.get("status") in {"paused", "queued"}), None)
        if task is None:
            self.queue_running = False
            self._render_table()
            if callable(self.busy_changed):
                self.busy_changed(False)
            self.progress_message.setText("Очередь завершена.")
            return
        if not Path(task["path"]).is_file():
            task["status"] = "failed"
            task["error"] = "Исходный файл больше не найден."
            task["finished_at"] = _now()
            self._save_queue()
            QTimer.singleShot(0, self._start_next)
            return

        resume_from_seconds = max(0.0, float(task.get("resume_seconds", 0) or 0))
        saved_prefix = str(task.get("partial_text", "") or "") if resume_from_seconds > 0 else ""
        was_resuming = bool(saved_prefix and resume_from_seconds > 0)
        if not was_resuming:
            resume_from_seconds = 0.0
            task["resume_seconds"] = 0.0
            task["progress"] = 0
            task["partial_text"] = ""
            task["segments"] = []

        task["run_elapsed_base"] = max(0, int(task.get("elapsed", 0) or 0))
        task["status"] = "processing"
        task["started_at"] = _now()
        task["started_epoch"] = time.time()
        task["error"] = ""
        self.active_task_id = task["id"]
        self._save_queue()
        self._render_table(task["id"])
        self.current_label.setText(Path(task["path"]).name)
        self.elapsed_label.setText("Прошло 00:00")
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(int(task.get("progress", 0) or 0))
        self.result_editor.setReadOnly(True)
        if self._displayed_task_id != task["id"]:
            self.result_editor.setPlainText(saved_prefix)
        self._displayed_task_id = task["id"]
        self.progress_message.setText(
            f"Продолжаем с {_format_duration(resume_from_seconds)}…" if resume_from_seconds > 0
            else "Подготовка аудио и загрузка модели…"
        )
        prefix_segments = list(task.get("segments", [])) if was_resuming else []
        self.active_worker = self.worker_factory(
            task["path"], resume_from_seconds=resume_from_seconds, prefix_text=saved_prefix,
            prefix_segments=prefix_segments
        )
        self.active_worker.progress_changed.connect(self._on_progress)
        if hasattr(self.active_worker, "segments_changed"):
            self.active_worker.segments_changed.connect(self._on_segments_changed)
        self.active_worker.succeeded.connect(self._on_success)
        self.active_worker.failed.connect(self._on_failure)
        self.active_worker.finished.connect(self._on_worker_finished)
        if callable(self.busy_changed):
            self.busy_changed(True)
        self.active_worker.start()
        self._render_table(task["id"])

    def _on_segments_changed(self, segments):
        task = self._task_by_id(self.active_task_id)
        if task is None:
            return
        # Store real Whisper timestamps as each segment is decoded so they survive a stop/crash.
        task["segments"] = [dict(item) for item in segments if isinstance(item, dict)]
        self._save_queue()

    def _result_editor_text_changed(self):
        if getattr(self, "_loading_editor", False):
            return
        task = self._task_by_id(self._displayed_task_id) if self._displayed_task_id else None
        if task is None or task.get("status") != "completed":
            return
        self._editor_save_timer.start()

    def _persist_result_editor(self):
        task = self._task_by_id(self._displayed_task_id) if self._displayed_task_id else None
        if task is None or task.get("status") != "completed":
            return
        text = self.result_editor.toPlainText()
        task["partial_text"] = text
        task["edited"] = True
        result_path = Path(task.get("result_path", "")) if task.get("result_path") else None
        try:
            if result_path:
                result_path.parent.mkdir(parents=True, exist_ok=True)
                result_path.write_text(text, encoding="utf-8", newline="")
            self._save_queue()
            self.progress_message.setText("Изменения текста автоматически сохранены.")
        except OSError as exc:
            self.progress_message.setText(f"Не удалось сохранить изменения: {exc}")

    def _on_progress(self, percent, message, partial_text, resume_seconds=0.0):
        task = self._task_by_id(self.active_task_id)
        if task is None:
            return
        if percent >= 0:
            task["progress"] = max(0, min(99, int(percent)))
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(task["progress"])
        elif "пауза" in (message or "").lower():
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(int(task.get("progress", 0) or 0))
        else:
            self.progress_bar.setRange(0, 0)
        self.progress_message.setText(message)
        if partial_text:
            task["partial_text"] = partial_text
        task["resume_seconds"] = max(
            float(task.get("resume_seconds", 0) or 0), float(resume_seconds or 0)
        )
        if self._displayed_task_id == task["id"]:
            if partial_text:
                self.result_editor.setPlainText(partial_text)
            self.result_editor.setReadOnly(True)
        self._save_queue()
        self._render_table()

    def _on_success(self, text, source_path):
        task = self._task_by_id(self.active_task_id)
        if task is None:
            return
        if self.active_worker is not None and hasattr(self.active_worker, "segments"):
            task["segments"] = [dict(item) for item in self.active_worker.segments]
        result_path = self.results_dir / f"{task['id']}.txt"
        temporary = result_path.with_suffix(".txt.tmp")
        temporary.write_text(text, encoding="utf-8", newline="")
        os.replace(temporary, result_path)
        task["status"] = "completed"
        task["progress"] = 100
        task["partial_text"] = text
        # This value is no longer used for completed jobs, but keep it monotonic if
        # waveform duration metadata hasn't been populated yet.
        task["resume_seconds"] = max(
            float(task.get("resume_seconds", 0) or 0),
            float(task.get("duration_ms", 0) or 0) / 1000.0,
        )
        task["result_path"] = str(result_path)
        task["finished_at"] = _now()
        started_epoch = float(task.get("started_epoch", time.time()) or time.time())
        task["elapsed"] = max(0, int(task.get("run_elapsed_base", 0) or 0)) + round(time.time() - started_epoch)
        task["error"] = ""
        self._save_queue()
        self.progress_bar.setRange(0, 100)
        self.progress_bar.setValue(100)
        self.progress_message.setText(f"Готово за {_format_duration(task['elapsed'])}")
        if self._displayed_task_id == task["id"]:
            self.result_editor.setReadOnly(False)
            self.result_editor.setPlainText(text)
        self._render_table()
        logging.getLogger("processing.queue").info(
            "Queue task completed: %s duration=%.2fs chars=%d", source_path, task["elapsed"], len(text)
        )

    def _on_failure(self, error_text, source_path):
        task = self._task_by_id(self.active_task_id)
        if task is None:
            return
        if self.active_worker is not None:
            if getattr(self.active_worker, "latest_text", ""):
                task["partial_text"] = str(self.active_worker.latest_text)
            task["resume_seconds"] = max(float(task.get("resume_seconds", 0) or 0), float(getattr(self.active_worker, "_checkpoint_seconds", 0) or 0))
            if hasattr(self.active_worker, "segments"):
                task["segments"] = [dict(item) for item in self.active_worker.segments]

        if self.shutdown_requested:
            resumable = bool(task.get("partial_text")) and float(task.get("resume_seconds", 0) or 0) > 0
            task["status"] = "paused" if resumable else "queued"
            task["error"] = "Сеанс остановлен при закрытии программы. Сохранённый прогресс будет восстановлен."
            task["finished_at"] = _now()
            self.progress_message.setText("Состояние сохранено для следующего запуска СПИКЕР.")
        elif self.discard_requested or self.cancel_requested:
            task["status"] = "cancelled"
            task["error"] = "Задача отменена пользователем."
            task["finished_at"] = _now()
            task["progress"] = 0
            task["resume_seconds"] = 0.0
            task["partial_text"] = ""
            task["segments"] = []
            task["elapsed"] = 0
            if self._displayed_task_id == task["id"]:
                self.result_editor.clear()
            self.progress_message.setText("Задача отменена. Её текст и контрольная точка сброшены.")
        elif self.stop_requested:
            task["status"] = "paused"
            task["error"] = "Остановлено с сохранением. Нажмите «Очередь», чтобы продолжить с контрольной точки."
            task["finished_at"] = _now()
            self.progress_message.setText(
                f"Остановлено с сохранением на {_format_duration(float(task.get('resume_seconds', 0) or 0))}. "
                "Можно продолжить."
            )
        else:
            was_interrupted = (
                "отменена пользователем" in (error_text or "").lower()
                or "остановлена пользователем" in (error_text or "").lower()
            )
            task["status"] = "paused" if was_interrupted and task.get("partial_text") else "failed"
            task["error"] = error_text
            task["finished_at"] = _now()
            if task["status"] == "failed":
                self.progress_message.setText(f"Ошибка: {error_text}")
                logging.getLogger("app.transcription_errors").error(
                    "Task failed: %s: %s", source_path, error_text
                )
            else:
                self.progress_message.setText("Работа остановлена. Сохранённые фрагменты можно продолжить.")

        self._save_queue()
        self._render_table()
        if self._displayed_task_id == task["id"]:
            self.result_editor.setReadOnly(task.get("status") == "processing")
            if task.get("partial_text"):
                self.result_editor.setPlainText(task["partial_text"])

    def _on_worker_finished(self):
        self.active_worker = None
        self.active_task_id = None
        self.cancel_requested = False
        self.stop_requested = False
        self.discard_requested = False
        self._render_table()
        if self.queue_running and any(task.get("status") in {"queued", "paused"} for task in self.tasks):
            QTimer.singleShot(0, self._start_next)
            return
        self.queue_running = False
        if callable(self.busy_changed):
            self.busy_changed(False)
        self._clock.start()

    def toggle_pause(self):
        worker = self.active_worker
        if worker is None:
            return
        if worker.is_paused():
            worker.resume()
            self.pause_button.setText("Пауза")
            self.progress_message.setText("Продолжаем распознавание…")
        else:
            worker.pause()
            self.pause_button.setText("Продолжить")
            self.progress_message.setText("Пауза запрошена; она вступит в силу на границе текущего сегмента.")

    def stop_current_task(self):
        """Stop the active worker, save its checkpoint and leave it resumable."""
        if self.active_worker is None:
            return
        self.queue_running = False
        self.stop_requested = True
        self.cancel_requested = False
        self.discard_requested = False
        self.progress_message.setText("Останавливаем на ближайшей границе сегмента и сохраняем прогресс…")
        self.active_worker.cancel()
        self.cancel_button.setEnabled(False)

    def cancel_selected_task(self):
        """Cancel a selected task and discard its saved checkpoint; never delete its source media."""
        task = self._selected_task()
        if not task:
            return
        if task.get("status") == "completed":
            QMessageBox.information(
                self, "Задача завершена",
                "Готовый результат можно удалить из списка отдельной кнопкой."
            )
            return
        if task.get("status") == "processing":
            if task.get("id") != self.active_task_id or self.active_worker is None:
                return
            answer = QMessageBox.question(
                self, "Отменить распознавание?",
                "Текущая задача будет остановлена, а сохранённый промежуточный текст и точка продолжения будут сброшены.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
            self.cancel_requested = True
            self.discard_requested = True
            self.stop_requested = False
            self.progress_message.setText("Отменяем задачу и сбрасываем её контрольную точку…")
            self.active_worker.cancel()
            self.cancel_task_button.setEnabled(False)
            return

        task["status"] = "cancelled"
        task["error"] = "Задача отменена пользователем."
        task["finished_at"] = _now()
        task["progress"] = 0
        task["resume_seconds"] = 0.0
        task["partial_text"] = ""
        task["elapsed"] = 0
        if self._displayed_task_id == task["id"]:
            self.result_editor.clear()
            self.result_editor.setReadOnly(True)
        self._save_queue()
        self._render_table(task["id"])
        self.progress_message.setText("Задача отменена. Исходный файл не удалён; можно повторить с начала.")

    def retry_selected(self):
        task = self._selected_task()
        if not task or task.get("status") not in {"failed", "cancelled", "paused"}:
            return
        if task.get("status") == "cancelled":
            task["progress"] = 0
            task["elapsed"] = 0
            task["resume_seconds"] = 0.0
            task["partial_text"] = ""
            if self._displayed_task_id == task["id"]:
                self.result_editor.clear()
                self.result_editor.setReadOnly(True)
        # Paused/failed tasks retain their committed transcript and checkpoint.
        task["status"] = "queued"
        task["error"] = ""
        task["started_at"] = ""
        task["finished_at"] = ""
        self._save_queue()
        self._render_table(task["id"])
        if float(task.get("resume_seconds", 0) or 0) > 0:
            self.progress_message.setText(
                f"Задача готова к продолжению с {_format_duration(float(task['resume_seconds']))}."
            )
        else:
            self.progress_message.setText("Задача возвращена в очередь с начала.")

    def remove_selected(self):
        task = self._selected_task()
        if not task:
            return
        if task.get("status") == "processing":
            QMessageBox.information(self, "Задача выполняется", "Сначала остановите текущую обработку.")
            return
        if task.get("status") == "completed":
            answer = QMessageBox.question(
                self, "Удалить из истории?",
                "Убрать задачу из списка? Уже сохранённый текстовый результат останется в папке results.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
                QMessageBox.StandardButton.No,
            )
            if answer != QMessageBox.StandardButton.Yes:
                return
        if self._waveform_worker is not None and self._waveform_task_id == task.get("id"):
            self._waveform_preempted_ids.discard(task.get("id"))
            self._waveform_worker.cancel()
        self._waveform_pending_ids = [task_id for task_id in self._waveform_pending_ids if task_id != task.get("id")]
        self._waveform_samples_by_path.pop(str(Path(task.get("path", "")).resolve()), None)
        self.tasks = [item for item in self.tasks if item.get("id") != task.get("id")]
        self._save_queue()
        self._render_table()
        self.result_editor.clear()
        self._displayed_task_id = None

    def _selection_changed(self):
        task = self._selected_task()
        if task is None:
            self._set_audio_source("")
            self.player_title.setText("Выберите файл")
            self.player_metadata.setText("Длительность · Размер · Формат")
            self.current_label.setText("Распознавание не запущено")
            self.elapsed_label.clear()
            self.result_editor.setReadOnly(True)
            if not self._loading_editor:
                self._loading_editor = True
                self.result_editor.clear()
                self._loading_editor = False
            return

        is_new_selection = task.get("id") != self._displayed_task_id
        if is_new_selection:
            if hasattr(self, "_editor_save_timer") and self._editor_save_timer.isActive():
                self._editor_save_timer.stop()
                self._persist_result_editor()
            self._displayed_task_id = task.get("id")
            self._set_audio_source(task.get("path", ""))
            self._loading_editor = True
            result_path = Path(task.get("result_path", "")) if task.get("result_path") else None
            if task.get("status") == "completed" and result_path and result_path.is_file():
                self.result_editor.setPlainText(result_path.read_text(encoding="utf-8"))
            else:
                self.result_editor.setPlainText(task.get("partial_text", ""))
            self.result_editor.setReadOnly(task.get("status") in {"queued", "processing"})
            self._loading_editor = False

        path = Path(task.get("path", ""))
        self.player_title.setText(path.name or "Выберите файл")
        duration_ms = int(task.get("duration_ms", 0) or 0)
        duration_text = _format_duration(duration_ms / 1000) if duration_ms else "Определяется…"
        self.player_metadata.setText(
            f"{duration_text}   ·   {_format_size(int(task.get('file_size', 0) or 0))}   ·   "
            f"{task.get('format') or path.suffix.upper().lstrip('.') or '—'}"
        )

        source = str(path.resolve()) if path else ""
        samples = self._waveform_samples_by_path.get(source)
        if samples and task.get("waveform_status") == "ready":
            self.waveform.set_samples(samples, duration_ms)
            self.waveform.set_loading(False)
        elif task.get("waveform_status") in {"queued", "loading"}:
            self.waveform.set_loading(True)
            if task.get("waveform_status") == "queued":
                self._queue_waveform_task(task.get("id"), priority=True)
        elif task.get("waveform_status") == "failed":
            self.waveform.clear()
            self.waveform.set_loading(False)

        is_active_task = task.get("id") == self.active_task_id and self.active_worker is not None
        no_active_task = self.active_worker is None
        if is_active_task or no_active_task:
            self.current_label.setText(path.name or "Распознавание не запущено")
        if is_active_task or no_active_task:
            if task.get("status") == "failed":
                self.progress_message.setText(task.get("error") or "Обработка завершилась с ошибкой.")
            elif task.get("status") == "queued" and task.get("error"):
                self.progress_message.setText(task.get("error"))
            self.progress_bar.setRange(0, 100)
            self.progress_bar.setValue(int(task.get("progress", 0) or 0))
            if no_active_task:
                self.elapsed_label.clear()

    def _update_clock(self):
        task = self._task_by_id(self.active_task_id)
        if task is not None and task.get("status") == "processing":
            run_seconds = max(0, int(time.time() - float(task.get("started_epoch", time.time()))))
            elapsed = max(0, int(task.get("run_elapsed_base", 0) or 0)) + run_seconds
            task["elapsed"] = elapsed
            paused = self.active_worker is not None and self.active_worker.is_paused()
            self.elapsed_label.setText(("Пауза запрошена · " if paused else "Прошло ") + _format_duration(elapsed))
            progress = int(task.get("progress", 0) or 0)
            if paused:
                self.progress_message.setText("Пауза запрошена. Текущий сегмент завершится, затем обработка остановится.")
            elif progress > 0:
                remaining = max(0, int(elapsed * (100 - progress) / progress))
                self.progress_message.setText(
                    f"Распознавание: {progress}% · Осталось примерно {_format_duration(remaining)}"
                )
            # Save time and progress periodically so the queue remains useful after a crash.
            if elapsed % 5 == 0:
                self._save_queue()
            row = next((i for i, t in enumerate(self.tasks) if t.get("id") == task["id"]), -1)
            if row >= 0 and self.table.item(row, 1):
                self.table.item(row, 1).setText(f"Обрабатывается · {_format_duration(elapsed)}")

    def _set_audio_source(self, raw_path: str):
        path = str(Path(raw_path).resolve()) if raw_path else ""
        if path == self._preview_source:
            task = self._task_for_source(path)
            if task:
                self._sync_player_enabled(task)
            return

        self.media_player.stop()
        self._preview_source = path
        self._player_error_sources.discard(path)
        self.waveform.clear()
        self.seek_slider.setRange(0, 0)
        self.audio_time_label.setText("00:00 / 00:00")
        self.player_controls.set_playing(False)
        task = self._task_for_source(path)
        if path and Path(path).is_file():
            self.audio_engine.load(path)
            self.audio_engine.set_rate(float(self.speed_combo.currentData() or 1.0))
            if task:
                duration_ms = int(task.get("duration_ms", 0) or 0)
                samples = self._waveform_samples_by_path.get(path)
                if samples:
                    self.waveform.set_samples(samples, duration_ms)
                elif task.get("waveform_status") in {"queued", "loading"}:
                    self.waveform.set_loading(True)
                self.player_metadata.setText(
                    f"{_format_duration(duration_ms / 1000) if duration_ms else 'Определяется…'}   ·   "
                    f"{_format_size(int(task.get('file_size', 0) or 0))}   ·   {task.get('format', '—')}"
                )
                if task.get("waveform_status") == "failed":
                    self.audio_status.setText(
                        task.get("waveform_error") or "Не удалось найти или прочитать аудиодорожку."
                    )
                elif task.get("waveform_status") in {"queued", "loading"}:
                    self.audio_status.setText("Подготавливаем аудиодорожку…")
                else:
                    self.audio_status.setText("Аудиодорожка готова. Можно прослушать файл и перейти по волне.")
                self._sync_player_enabled(task)
            else:
                self.player_controls.set_enabled(False)
                self.audio_status.setText("Загружаем сведения об аудиодорожке…")
            # QMediaPlayer and FFmpeg are independent: let both inspect the selected source.
        else:
            self.audio_engine.load("")
            self.player_controls.set_enabled(False)
            self.waveform_button.setEnabled(False)
            if path:
                self.audio_status.setText("Исходный файл не найден. Добавьте его в очередь заново.")
            else:
                self.audio_status.setText("Выберите файл из очереди.")
            self.player_metadata.setText("Длительность · Размер · Формат")
        if path and Path(path).is_file():
            self.waveform_button.setEnabled(bool(task and task.get("waveform_status") != "loading"))

    def _task_for_source(self, path: str):
        if not path:
            return None
        key = MediaQueue.path_key(path)
        return next((item for item in self.tasks
                     if item.get("path") and MediaQueue.path_key(item["path"]) == key), None)

    def _sync_player_enabled(self, task):
        ready = bool(
            task and task.get("audio_stream_available") is True
            and self._preview_source not in self._player_error_sources
        )
        self.player_controls.set_enabled(ready)
        seekable = bool(ready and self.media_player.isSeekable())
        for widget in (
            self.player_controls.back10_button, self.player_controls.forward10_button,
            self.player_controls.seek_slider,
        ):
            widget.setEnabled(seekable)
        self.waveform_button.setEnabled(bool(task and Path(task.get("path", "")).is_file()
                                             and task.get("waveform_status") != "loading"))
        if not ready:
            self.player_controls.set_playing(False)

    def toggle_preview(self):
        task = self._task_for_source(self._preview_source)
        if not task or task.get("audio_stream_available") is not True:
            self.audio_status.setText("Сначала дождитесь проверки аудиодорожки.")
            return
        self.audio_engine.play_pause()

    def stop_preview(self):
        self.audio_engine.stop()
        self._preview_position_changed(0)

    def _step_preview(self, delta_ms: int):
        if not self.media_player.isSeekable():
            self.audio_status.setText("Для этого файла перемотка пока недоступна.")
            return
        if not self.audio_engine.step(delta_ms):
            self.audio_status.setText("Для этого файла перемотка не поддерживается.")

    def _seek_preview(self, position_ms: int):
        if not self.media_player.isSeekable():
            self.audio_status.setText("Для этого файла перемотка пока недоступна: медиадвижок ещё не открыл временную шкалу.")
            return
        task = self._task_for_source(self._preview_source)
        target_ms = int(position_ms)
        waveform_duration = int(task.get("duration_ms", 0) or 0) if task else 0
        player_duration = int(self.media_player.duration() or 0)
        if waveform_duration > 0 and player_duration > 0 and waveform_duration != player_duration:
            target_ms = int(target_ms * player_duration / waveform_duration)
        if not self.audio_engine.seek(target_ms):
            self.audio_status.setText("Для этого файла перемотка не поддерживается.")

    def _set_volume(self, value: int):
        self.audio_engine.set_volume(value)

    def _set_playback_rate(self, rate: float):
        if not self.audio_engine.set_rate(rate):
            self.audio_status.setText("Изменение скорости не поддерживается этим медиадвижком.")

    def _preview_position_changed(self, position_ms: int):
        if not self.seek_slider.isSliderDown():
            self.seek_slider.setValue(int(position_ms))
        duration = int(self.media_player.duration() or 0)
        task = self._task_for_source(self._preview_source)
        waveform_duration = (
            int(task.get("duration_ms", 0) or 0)
            if task and task.get("waveform_status") == "ready"
            else duration
        )
        self.waveform.set_playhead(position_ms, waveform_duration)
        self._update_audio_time(position_ms, duration)

    def _preview_duration_changed(self, duration_ms: int):
        duration_ms = max(0, int(duration_ms))
        self.seek_slider.setRange(0, duration_ms)
        if not self.seek_slider.isSliderDown():
            self.seek_slider.setValue(min(int(self.media_player.position()), duration_ms))
        # Qt's duration is useful for the player. FFmpeg's decoded audio duration
        # remains authoritative for the waveform timeline when available.
        task = self._task_for_source(self._preview_source)
        if task and duration_ms and task.get("waveform_status") != "ready":
            task["player_duration_ms"] = duration_ms
            if not task.get("duration_ms"):
                task["duration_ms"] = duration_ms
            self.player_metadata.setText(
                f"{_format_duration(duration_ms / 1000)}   ·   "
                f"{_format_size(int(task.get('file_size', 0) or 0))}   ·   {task.get('format', '—')}"
            )
            self._save_queue()
            self._render_table()
        task = self._task_for_source(self._preview_source)
        waveform_duration = (
            int(task.get("duration_ms", 0) or 0)
            if task and task.get("waveform_status") == "ready"
            else duration_ms
        )
        self.waveform.set_playhead(self.media_player.position(), waveform_duration)
        self._update_audio_time(self.media_player.position(), duration_ms)

    def _update_audio_time(self, pos_ms: int, duration_ms: int):
        self.audio_time_label.setText(
            f"{_format_duration(pos_ms / 1000)} / {_format_duration(duration_ms / 1000)}"
        )

    def _preview_state_changed(self, state):
        playing = state == QMediaPlayer.PlaybackState.PlayingState
        self.player_controls.set_playing(playing)
        if not playing and self.media_player.mediaStatus() == QMediaPlayer.MediaStatus.EndOfMedia:
            self._preview_position_changed(self.media_player.duration())

    def _preview_media_status_changed(self, status):
        if status == QMediaPlayer.MediaStatus.LoadingMedia:
            self.audio_status.setText("Медиадвижок открывает файл…")
        elif status == QMediaPlayer.MediaStatus.LoadedMedia:
            self.audio_engine.set_rate(float(self.speed_combo.currentData() or 1.0))
            self.audio_status.setText("Файл загружен в медиадвижок.")
            self._preview_duration_changed(self.media_player.duration())
            self._preview_seekable_changed(self.media_player.isSeekable())
            self._sync_player_enabled(self._task_for_source(self._preview_source))
        elif status == QMediaPlayer.MediaStatus.EndOfMedia:
            self.player_controls.set_playing(False)
            self._preview_position_changed(self.media_player.duration())
        elif status == QMediaPlayer.MediaStatus.InvalidMedia:
            self._player_error_sources.add(self._preview_source)
            self.audio_status.setText(
                "Windows/Qt не смог воспроизвести этот формат или кодек. Транскрибация через FFmpeg может оставаться доступной."
            )
            self.player_controls.set_enabled(False)
            self.player_controls.set_playing(False)

    def _preview_seekable_changed(self, seekable):
        task = self._task_for_source(self._preview_source)
        ready = bool(task and task.get("audio_stream_available") is True)
        for button in (self.player_controls.back10_button, self.player_controls.forward10_button):
            button.setEnabled(bool(seekable and ready))
        self.seek_slider.setEnabled(bool(seekable and ready))

    def _preview_error(self, error, message):
        if self._preview_source:
            self._player_error_sources.add(self._preview_source)
        if message:
            self.audio_status.setText(
                f"Не удалось воспроизвести этот файл: {message}. Проверьте поддержку кодека Windows/Qt."
            )
        self.player_controls.set_enabled(False)
        self.player_controls.set_playing(False)

    def build_waveform(self, force=False):
        task = self._task_for_source(self._preview_source)
        if not task or not Path(self._preview_source).is_file():
            return
        if self._waveform_worker is not None and self._waveform_worker.isRunning():
            if self._waveform_task_id == task.get("id"):
                return
            self._queue_waveform_task(task.get("id"), priority=True, force=bool(force))
            self.audio_status.setText("Текущий анализ завершится первым; выбранный файл следующий в очереди анализа.")
            return
        self.audio_status.setText("Подготавливаем аудиодорожку…")
        self.waveform.set_loading(True)
        self._queue_waveform_task(task.get("id"), priority=True, force=bool(force))

    def _audio_stream_detected(self, source_path: str):
        task = self._task_for_source(source_path)
        if task is None:
            return
        task["audio_stream_available"] = True
        self._save_queue()
        if str(Path(source_path).resolve()) == self._preview_source:
            self.audio_status.setText("Аудиодорожка найдена. Подготавливаем звуковую волну…")
            self._sync_player_enabled(task)
            self.waveform.set_playhead(self.media_player.position(), self.media_player.duration())

    def _waveform_ready(self, source_path: str, bars, duration_ms: int, cache_key: str):
        task = self._task_for_source(source_path)
        if task is None:
            return
        task["duration_ms"] = int(duration_ms)
        task["audio_stream_available"] = True
        task["waveform_status"] = "ready"
        task["waveform_cache_key"] = cache_key
        task["waveform_error"] = ""
        resolved = str(Path(source_path).resolve())
        self._waveform_samples_by_path[resolved] = list(bars)
        self._save_queue()
        self._render_table()
        if resolved == self._preview_source:
            self.waveform.set_samples(bars, duration_ms)
            self.waveform.set_playhead(self.media_player.position(), duration_ms)
            self.player_metadata.setText(
                f"{_format_duration(duration_ms / 1000)}   ·   "
                f"{_format_size(int(task.get('file_size', 0) or 0))}   ·   {task.get('format', '—')}"
            )
            if resolved in self._player_error_sources:
                self.audio_status.setText(
                    "Аудиодорожка найдена, но Qt/Windows не смог воспроизвести этот кодек. "
                    "Попробуйте другой формат; транскрибация через FFmpeg может работать."
                )
            else:
                self.audio_status.setText(
                    "Аудиодорожка готова. Нажмите «Воспроизвести» или выберите позицию на волне."
                )
            self._sync_player_enabled(task)
            self._update_audio_time(self.media_player.position(), max(duration_ms, self.media_player.duration()))
        self.waveform_button.setEnabled(True)

    def _waveform_failed(self, source_path: str, error_text: str):
        task = self._task_for_source(source_path)
        if task is not None:
            task["waveform_status"] = "failed"
            task["audio_stream_available"] = False
            task["waveform_error"] = error_text
            task["duration_ms"] = 0
            self._save_queue()
            self._render_table()
        if str(Path(source_path).resolve()) == self._preview_source:
            self.waveform.clear()
            self.waveform.set_loading(False)
            self.player_controls.set_enabled(False)
            self.waveform_button.setEnabled(True)
            lowered = error_text.lower()
            if any(term in lowered for term in ("аудиодорожка", "не найдено аудиоданных", "matches no streams", "does not contain any stream")):
                message = "В файле не найдена читаемая аудиодорожка. Проверьте, что это аудио или видео со звуком."
            elif any(term in lowered for term in ("invalid data", "moov atom not found", "error opening input", "corrupt", "could not find codec")):
                message = "Файл повреждён или его кодек не поддерживается встроенным декодером."
            else:
                message = "Не удалось прочитать аудио. Проверьте файл и попробуйте другой формат."
            task = self._task_for_source(source_path)
            if task:
                self.player_metadata.setText(
                    f"Длительность не определена   ·   {_format_size(int(task.get('file_size', 0) or 0))}   ·   "
                    f"{task.get('format', '—')}"
                )
            self.audio_status.setText(message)

    def _waveform_worker_finished(self):
        worker = self._waveform_worker
        task_id = self._waveform_task_id
        self._waveform_worker = None
        self._waveform_task_id = None
        task = self._task_by_id(task_id) if task_id else None
        preempted = bool(task_id and task_id in self._waveform_preempted_ids)
        self._waveform_preempted_ids.discard(task_id)
        if task and task.get("waveform_status") == "loading":
            # If a different file was selected, resume the old analysis later instead of
            # making the new selection wait behind a large video.
            task["waveform_status"] = "queued"
            if task_id not in self._waveform_pending_ids:
                if preempted:
                    self._waveform_pending_ids.append(task_id)
                else:
                    self._waveform_pending_ids.insert(0, task_id)
        self._save_queue()
        QTimer.singleShot(0, self._start_next_waveform_extraction)

    def send_result_to_ai(self):
        task = self._selected_task()
        text = self.result_editor.toPlainText().strip()
        if not task or not text:
            QMessageBox.information(self, "Нет расшифровки", "Выберите задачу и убедитесь, что её результат не пуст.")
            return
        self.ai_requested.emit(text, str(task.get("path", "")))

    @staticmethod
    def _subtitle_time(seconds, vtt=False):
        from subtitle_export import format_timestamp
        return format_timestamp(seconds, vtt)

    def _render_subtitles(self, segments, vtt=False):
        return _render_subtitles(segments, vtt)

    def save_current_result_as(self):
        task = self._selected_task()
        text = self.result_editor.toPlainText()
        if not task or not text.strip():
            self.progress_message.setText("Нет текста для сохранения.")
            return
        default_name = (Path(task["path"]).stem or "Spiker_result") + ".txt"
        filters = "Текстовый файл (*.txt);;Markdown (*.md);;Субтитры SubRip (*.srt);;WebVTT (*.vtt);;JSON с сегментами (*.json);;Документ Word (*.docx)"
        target, selected_filter = QFileDialog.getSaveFileName(self, "Экспорт результата", default_name, filters)
        if not target:
            return
        extensions = {
            "Текстовый файл": ".txt", "Markdown": ".md", "SubRip": ".srt",
            "WebVTT": ".vtt", "JSON": ".json", "Word": ".docx",
        }
        desired_extension = next((ext for label, ext in extensions.items() if label in selected_filter), ".txt")
        path = Path(target)
        if path.suffix.lower() not in set(extensions.values()):
            path = path.with_suffix(desired_extension)
        elif desired_extension != ".txt" and path.suffix.lower() != desired_extension:
            path = path.with_suffix(desired_extension)
        try:
            suffix = path.suffix.lower()
            segments = task.get("segments", [])
            if suffix in {".srt", ".vtt"}:
                output = self._render_subtitles(segments, vtt=(suffix == ".vtt"))
                if not output:
                    QMessageBox.warning(self, "Нет временных отметок", "Для этой расшифровки нет сегментов с временными отметками. Субтитры нельзя создать достоверно. Перезапустите распознавание в версии с поддержкой временных сегментов.")
                    return
                path.write_text(output, encoding="utf-8-sig" if suffix == ".srt" else "utf-8", newline="")
            elif suffix == ".json":
                payload = {"source_path": task.get("path", ""), "text": text,
                           "segments": segments, "created_at": _now(), "edited": bool(task.get("edited", False))}
                path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
            elif suffix == ".docx":
                try:
                    from docx import Document
                except ImportError as exc:
                    raise RuntimeError("Для экспорта DOCX установите зависимость python-docx: python -m pip install python-docx.") from exc
                document = Document()
                document.add_heading(Path(task.get("path", "Расшифровка")).stem, level=1)
                for paragraph in text.split("\n"):
                    document.add_paragraph(paragraph)
                document.save(str(path))
            elif suffix == ".md":
                path.write_text(f"# {Path(task.get('path', 'Расшифровка')).stem}\n\n{text}\n", encoding="utf-8", newline="")
            else:
                path.write_text(text, encoding="utf-8", newline="")
            # The canonical TXT result and queue copy always track edits, independent of export format.
            task["partial_text"] = text
            if task.get("result_path"):
                Path(task["result_path"]).write_text(text, encoding="utf-8", newline="")
            self._save_queue()
            self.progress_message.setText(f"Результат сохранён: {path}")
        except Exception as exc:
            QMessageBox.warning(self, "Ошибка сохранения", str(exc))

    def open_selected_source(self, row=None, column=None):
        task = self._selected_task()
        if task and Path(task.get("path", "")).is_file():
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(Path(task["path"]).resolve())))

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls() and any(url.isLocalFile() for url in event.mimeData().urls()):
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        paths = [url.toLocalFile() for url in event.mimeData().urls() if url.isLocalFile()]
        self.add_paths(paths)
        event.acceptProposedAction()

    def prepare_for_shutdown(self, timeout_ms=12000) -> bool:
        """Persist the exact last checkpoint before asking native inference to stop."""
        if hasattr(self, "_editor_save_timer") and self._editor_save_timer.isActive():
            self._editor_save_timer.stop()
            self._persist_result_editor()
        self.media_player.stop()
        if self._waveform_worker is not None and self._waveform_worker.isRunning():
            self._waveform_worker.cancel()
            if not self._waveform_worker.wait(timeout_ms):
                return False
        if self.active_worker is None:
            self._save_queue()
            return True
        task = self._task_by_id(self.active_task_id)
        if task is not None:
            resumable = bool(task.get("partial_text")) and float(task.get("resume_seconds", 0) or 0) > 0
            task["status"] = "paused" if resumable else "queued"
            task["error"] = "Сеанс остановлен при закрытии программы. Сохранённый прогресс будет восстановлен."
            task["finished_at"] = _now()
            # Keep partial_text, segments, progress and resume_seconds intact on disk.
            self._save_queue()
        self.queue_running = False
        self.shutdown_requested = True
        self.cancel_requested = False
        self.stop_requested = False
        self.discard_requested = False
        worker = self.active_worker
        worker.cancel()
        if worker.wait(timeout_ms):
            # The GUI thread was waiting, so queued Qt signals may not yet have updated
            # the JSON snapshot. Read the final worker-owned checkpoint directly.
            if task is not None:
                latest_text = str(getattr(worker, "latest_text", "") or task.get("partial_text", ""))
                checkpoint = max(float(task.get("resume_seconds", 0) or 0), float(getattr(worker, "_checkpoint_seconds", 0) or 0))
                segments = [dict(item) for item in getattr(worker, "segments", task.get("segments", [])) if isinstance(item, dict)]
                if latest_text:
                    task["partial_text"] = latest_text
                task["resume_seconds"] = checkpoint
                task["segments"] = segments
                task["progress"] = max(int(task.get("progress", 0) or 0), min(99, int(checkpoint / max(float(task.get("duration_ms", 0) or 0) / 1000, 1) * 99))) if checkpoint > 0 else int(task.get("progress", 0) or 0)
                task["status"] = "paused" if latest_text and checkpoint > 0 else "queued"
                self._save_queue()
            self.active_worker = None
            self.active_task_id = None
            return True
        return False

